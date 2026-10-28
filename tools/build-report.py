from pathlib import Path
from lxml import html
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
import re
root=Path.cwd();tree=html.fromstring((root/'docs/development-report.html').read_text(encoding='utf-8'));d=Document();sec=d.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=Mm(16);sec.bottom_margin=Mm(16);sec.left_margin=Mm(18);sec.right_margin=Mm(18)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','List Bullet','List Number']:
 st=d.styles[name];st.font.name='맑은 고딕';st.font.color.rgb=RGBColor(0,0,0);st.font.italic=False;st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'맑은 고딕')
 st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.15
normal=d.styles['Normal'];normal.font.size=Pt(9.5)
for name,size in [('Title',28),('Subtitle',15),('Heading 1',17),('Heading 2',11),('Caption',8.5)]:d.styles[name].font.size=Pt(size)
d.styles['Title'].font.name='Georgia';d.styles['Title'].element.rPr.rFonts.set(qn('w:eastAsia'),'바탕')
for name in ['Heading 1','Heading 2']:
 d.styles[name].paragraph_format.space_before=Pt(9);d.styles[name].paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('최근일 2026505521   ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);footer.add_run(' / ');total=OxmlElement('w:fldSimple');total.set(qn('w:instr'),'NUMPAGES');footer._p.append(total)
replacements={'Agent':'Codex','fixture':'시험용 입력','사용자 UAT':'사용자 기능 확인','사용자의 직접 UAT':'사용자의 직접 기능 확인','현업 UAT':'실무 기능 확인','UAT':'사용자 기능 확인','스냅숏':'기준 정보','반복 결과의 일치':'반복 실행 결과의 일치','본 재무제표':'본 재무제표'}
def clean(t):
 for a,b in replacements.items():t=t.replace(a,b)
 return re.sub(r'\s+',' ',t).strip()
def txt(e):
 s=''.join(e.itertext());return clean(s)
def fill(p,e):
 def visit(node,bold=False):
  if node.text:
   r=p.add_run(re.sub(r'\s+',' ',node.text));r.bold=bold
  for child in node:
   if child.tag=='br':p.add_run().add_break()
   elif child.tag=='a' and child.get('href','').startswith('https://'):
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(child.get('href'),RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);text=OxmlElement('w:t');text.text=txt(child);r.append(text);link.append(r);p._p.append(link)
   else:visit(child,bold or child.tag in ['b','strong'])
   if child.tail:p.add_run(re.sub(r'\s+',' ',child.tail))
 visit(e)
 for run in p.runs:
  for a,b in replacements.items():run.text=run.text.replace(a,b)
 return p

def paragraph(e,style=None):
 return fill(d.add_paragraph(style=style),e)

def table(e):
 rows=e.xpath('./thead/tr|./tbody/tr|./tr'); cols=max(len(r.xpath('./th|./td')) for r in rows)
 t=d.add_table(rows=0,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths={2:[.28,.72],3:[.28,.20,.52],4:[.38,.20,.20,.22],6:[.32,.13,.15,.15,.12,.13]}.get(cols,[1/cols]*cols)
 for row_idx,row in enumerate(rows):
  cells=t.add_row().cells
  for i,(cell,el) in enumerate(zip(cells,row.xpath('./th|./td'))):
   cell.width=Mm(174*widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cell.text='';fill(cell.paragraphs[0],el)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.15
    if i>0 and re.fullmatch(r'[\d /]+',cell.text):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(8.6);r.bold=row_idx==0
   pr=cell._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
   pr.append(borders)
   if row_idx==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F5F9');pr.append(sh)
  if row_idx==0:
   repeat=OxmlElement('w:tblHeader');cells[0]._tc.getparent().get_or_add_trPr().append(repeat)
 d.add_paragraph().paragraph_format.space_after=Pt(1)

def figure(e):
 for img in e.xpath('.//img'):
  path=(root/'docs'/img.attrib['src']).resolve()
  if path.suffix=='.svg':continue
  with Image.open(path) as im:w,h=im.size
  width=min(174,78*w/h);p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Mm(width))
 for c in e.xpath('./figcaption'):paragraph(c,'Caption')

def emit(e):
 cls=e.attrib.get('class','')
 if 'folio' in cls or e.tag=='span':return
 if e.tag=='h1':
  d.add_paragraph('DART Filing Readiness',style='Title')
  for sub in e.xpath('./span'):paragraph(sub,'Subtitle')
 elif e.tag=='h2':paragraph(e,'Heading 1')
 elif e.tag=='h3':paragraph(e,'Heading 2')
 elif e.tag=='p':paragraph(e)
 elif e.tag=='table':table(e)
 elif e.tag=='figure':figure(e)
 elif e.tag in ['ul','ol']:
  for li in e:paragraph(li,'List Bullet' if e.tag=='ul' else 'List Number')
 elif e.tag=='div' and 'image-pair' in cls:
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for j,img in enumerate(e.xpath('.//img')):
   path=(root/'docs'/img.attrib['src']).resolve()
   with Image.open(path) as im:w,h=im.size
   width=min(112 if j==0 else 52,80*w/h)
   p.add_run().add_picture(str(path),width=Mm(width));p.add_run('  ')
  for cap in e.xpath('.//figcaption'):paragraph(cap,'Caption')
 elif e.tag=='div' and 'report-brand' in cls:
  d.add_paragraph('생성형 AI와 바이브 코딩 과제 2',style='Subtitle')
 elif e.tag=='div' and any(x in cls for x in ['author','meta','box','status','workflow']):paragraph(e)
 elif e.tag=='div' and 'flow' in cls:
  for el in e:paragraph(el)
 else:
  for el in e:emit(el)
pages=tree.xpath('//main/section')
for i,page in enumerate(pages):
 if i:d.add_page_break()
 for e in page:emit(e)
d.core_properties.title='공시 전 문서 점검 웹 애플리케이션 개발 보고서';d.core_properties.author='최근일';d.core_properties.subject='생성형 AI와 바이브 코딩 과제 2'
out=root/'submission/과제2_공시체크리스트_제출보고서.docx';# Remove inherited Word template title rules and theme decoration.
for part in [d.styles.element,d.element]:
 for node in list(part.iter(qn('w:pBdr'))):node.getparent().remove(node)
d.save(out);print(out)
